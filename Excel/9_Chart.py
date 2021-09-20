from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference, LineChart
wb = load_workbook("C:/Users/parkj/Desktop/프로그래밍/python/RPA/Excel/File/Sample.xlsx")
ws = wb.active

# B2 ~ C11 까지의 데이터를 차트로 생성
barValue = Reference(ws, min_row=2, max_row=11, min_col=2, max_col=3)
barChart = BarChart() # 차트 종류 설정(Bar, Line, Pie...)
barChart.add_data(barValue) # 차트 데이터 추가
ws.add_chart(barChart, "E1") # 넣을 차트와 위치 정의

# B1 ~ C11 까지의 데이터를 차트로 생성
lineValue = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3)
lineChart = LineChart()
lineChart.add_data(lineValue, titles_from_data = True) # 계열에 1번째 row값 대입
lineChart.title = "성적표" # 제목 정의
lineChart.style = 10 # 미리 정의된 스타일을 적용
lineChart.y_axis.title = "점수" # y축 제목 정의
lineChart.x_axis.title = "번호" # x축 제목 정의
ws.add_chart(lineChart, "E15")

wb.save("C:/Users/parkj/Desktop/프로그래밍/python/RPA/Excel/File/Sample_Chart.xlsx")
